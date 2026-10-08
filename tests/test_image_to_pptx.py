"""Image-to-PPTX conversion test: runs the real example pipeline and checks the output."""
import unittest
from pathlib import Path

from _common import REPO_ROOT
from convert_image_to_pptx import audit, build
from figure_kit import load_theme
from validate_pptx import validate
import yaml


class ImageToPptxTest(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.image = REPO_ROOT / "examples" / "image_to_pptx_input.png"
        cls.regions_file = REPO_ROOT / "examples" / "image_to_pptx_regions.yaml"
        cls.out = REPO_ROOT / "output" / "tests" / "test_image_to_pptx.pptx"

    def test_region_file_passes_audit(self):
        spec = yaml.safe_load(self.regions_file.read_text(encoding="utf-8"))
        self.assertEqual(audit(spec["regions"]), [], "region file has blocking issues")

    def test_conversion_output(self):
        spec = yaml.safe_load(self.regions_file.read_text(encoding="utf-8"))
        result = build(self.image, spec["regions"], self.out, load_theme())

        # PPTX exists and opens; editability report numbers are consistent
        self.assertTrue(self.out.exists())
        self.assertEqual(len(result["rebuild"]), 9)
        self.assertEqual(len(result["embed"]), 1)
        self.assertAlmostEqual(result["score"], 0.9)

        report = validate(self.out)
        # native: title + 4 nodes + legend; connectors: 3 arrows, all arrowed
        self.assertGreaterEqual(report["native_shapes"], 6)
        self.assertEqual(report["connectors"], 3)
        self.assertEqual(report["arrowed_connectors"], 3)
        # exactly one embedded picture (the sample image), NOT a full-slide screenshot
        self.assertEqual(report["pictures"], 1)
        self.assertFalse(any("screenshot" in w for w in report["warnings"]), report["warnings"])
        # Chinese labels survived into the PPTX
        from pptx import Presentation
        prs = Presentation(self.out)
        texts = " ".join(
            shape.text for slide in prs.slides
            for shape in slide.shapes if getattr(shape, "has_text_frame", False)
        )
        self.assertIn("肿瘤影像分析流程", texts)
        self.assertIn("特征提取", texts)


if __name__ == "__main__":
    unittest.main()
