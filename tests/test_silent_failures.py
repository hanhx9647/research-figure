"""Regression tests for the silent failure modes this skill exists to prevent.

Every case here produces **no exception** during generation. The file is only
wrong when you open it. If these tests ever stop failing for the right reasons,
the validator has gone blind again.
"""
import sys
import unittest
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))

from _common import OUTPUT  # noqa: E402
from figure_kit import add_arrow, add_node, hex_to_rgb, load_theme, new_presentation, set_font  # noqa: E402
from pptx.enum.shapes import MSO_SHAPE  # noqa: E402
from pptx.util import Inches  # noqa: E402
from validate_pptx import validate  # noqa: E402


class SilentFailureTest(unittest.TestCase):
    def test_unset_text_colour_is_detected(self):
        """A shape with no explicit text colour renders white — must be flagged."""
        theme = load_theme()
        prs, slide = new_presentation(theme)
        box = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE,
                                     Inches(1), Inches(1), Inches(2.4), Inches(0.8))
        box.fill.solid()
        box.fill.fore_color.rgb = hex_to_rgb("E2F0D9")
        tf = box.text_frame
        tf.clear()
        run = tf.paragraphs[0].add_run()
        run.text = "colour never set"
        set_font(run, theme["fonts"]["latin"], theme["fonts"]["cjk"], size=14)  # no colour!

        path = OUTPUT / "silent_unset_colour.pptx"
        path.parent.mkdir(parents=True, exist_ok=True)
        prs.save(path)

        report = validate(path)
        self.assertTrue(
            any("low text contrast" in w for w in report["warnings"]),
            f"validator missed invisible white text: {report}",
        )

    def test_theme_coloured_text_passes(self):
        theme = load_theme()
        prs, slide = new_presentation(theme)
        add_node(slide, 1.0, 1.0, "visible label 可见标签", theme["colors"]["process"], theme)
        path = OUTPUT / "silent_coloured_ok.pptx"
        path.parent.mkdir(parents=True, exist_ok=True)
        prs.save(path)

        report = validate(path)
        self.assertFalse(
            any("low text contrast" in w for w in report["warnings"]),
            report["warnings"],
        )

    def test_connector_without_arrowhead_is_detected(self):
        """`line.end_arrowhead` does not exist in python-pptx; a naive connector
        silently has no arrowhead at all."""
        from pptx.enum.shapes import MSO_CONNECTOR

        theme = load_theme()
        prs, slide = new_presentation(theme)
        conn = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT,
                                          Inches(1), Inches(1), Inches(3), Inches(1))
        conn.line.color.rgb = hex_to_rgb(theme["colors"]["connector"])
        path = OUTPUT / "silent_no_arrowhead.pptx"
        path.parent.mkdir(parents=True, exist_ok=True)
        prs.save(path)

        report = validate(path)
        self.assertEqual(report["connectors"], 1)
        self.assertEqual(report["arrowed_connectors"], 0)
        self.assertTrue(any("without arrowhead" in w for w in report["warnings"]), report)

    def test_arrowed_connector_passes(self):
        theme = load_theme()
        prs, slide = new_presentation(theme)
        add_arrow(slide, 1, 1, 3, 1, color=theme["colors"]["connector"])
        path = OUTPUT / "silent_arrow_ok.pptx"
        path.parent.mkdir(parents=True, exist_ok=True)
        prs.save(path)

        report = validate(path)
        self.assertEqual(report["arrowed_connectors"], report["connectors"])
        self.assertFalse(any("without arrowhead" in w for w in report["warnings"]), report)


if __name__ == "__main__":
    unittest.main(verbosity=2)
