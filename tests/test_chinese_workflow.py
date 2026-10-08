"""中文流程图测试：生成 → 校验 → 断言可编辑性与中文渲染。"""
import unittest

from _common import OUTPUT, assert_editable, build_chain
from pptx import Presentation
from validate_pptx import validate


class ChineseWorkflowTest(unittest.TestCase):
    def test_chinese_workflow(self):
        out = build_chain(
            [
                ("数据采集", "DDEBF7"),
                ("数据清洗", "E2F0D9"),
                ("模型训练", "FCE4D6"),
                ("结论输出", "FFE699"),
            ],
            OUTPUT / "test_chinese_workflow.pptx",
            feedback=(0, 2, False),   # 模型训练 -> 数据采集 反馈回路
        )
        self.assertTrue(out.exists())

        # PPTX 能否打开 + 中文能否读回
        prs = Presentation(out)
        texts = " ".join(
            shape.text for slide in prs.slides
            for shape in slide.shapes if getattr(shape, "has_text_frame", False)
        )
        for expected in ("数据采集", "数据清洗", "模型训练", "结论输出"):
            self.assertIn(expected, texts)

        report = validate(out)
        assert_editable(report, min_shapes=4, min_connectors=4, min_text=4)

        # 中文必须有 a:ea 字体（否则回退默认宋体）
        self.assertFalse(
            any("a:ea" in w for w in report["warnings"]),
            f"CJK font fallback missing: {report['warnings']}",
        )


if __name__ == "__main__":
    unittest.main()
